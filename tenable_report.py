"""Analyse l'évolution de rapports CSV Tenable et produit un rapport HTML autonome."""

from __future__ import annotations

import argparse
import base64
import csv
import html
import ipaddress
import io
import json
import sys
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator


FindingKey = tuple[str, str, str, str]
REQUIRED_COLUMNS = ("ip", "port", "plugin_id", "cve", "severity")
UNASSIGNED = "Non défini"


class ReportError(ValueError):
    """Erreur de configuration ou de données empêchant la génération du rapport."""


@dataclass(frozen=True, slots=True)
class Finding:
    ip: str
    port: str
    plugin_id: str
    cve: str
    site: str
    owner: str
    severity: str = UNASSIGNED

    @property
    def key(self) -> FindingKey:
        return (self.ip, self.port, self.plugin_id, self.cve)


@dataclass(frozen=True, slots=True)
class ScanReport:
    label: str
    source: Path
    findings: dict[FindingKey, Finding]


@dataclass(frozen=True, slots=True)
class Comparison:
    label: str
    new: tuple[Finding, ...]
    persistent: tuple[Finding, ...]
    fixed: tuple[Finding, ...]


def load_configuration(config_path: Path) -> tuple[list[ScanReport], Path]:
    """Charge et valide la configuration JSON, puis résout les chemins relatifs."""
    try:
        config = json.loads(config_path.read_text(encoding="utf-8-sig"))
    except OSError as exc:
        raise ReportError(f"Impossible de lire le fichier de configuration « {config_path} » : {exc}") from exc
    except UnicodeError as exc:
        raise ReportError(f"Le fichier de configuration « {config_path} » n'est pas encodé en UTF-8.") from exc
    except json.JSONDecodeError as exc:
        raise ReportError(f"Configuration JSON invalide à la ligne {exc.lineno} : {exc.msg}") from exc

    if not isinstance(config, dict):
        raise ReportError("La configuration doit être un objet JSON.")
    columns = config.get("columns")
    if not isinstance(columns, dict) or any(
        not isinstance(columns.get(name), str) or not columns[name].strip()
        for name in REQUIRED_COLUMNS
    ):
        raise ReportError(f"« columns » doit définir des en-têtes non vides pour : {', '.join(REQUIRED_COLUMNS)}.")
    for name in ("site", "owner"):
        if name in columns and columns[name] is not None and (
            not isinstance(columns[name], str) or not columns[name].strip()
        ):
            raise ReportError(f"L'en-tête « columns.{name} » doit être une chaîne non vide ou null.")
    mapped_headers = [columns[name] for name in REQUIRED_COLUMNS]
    mapped_headers.extend(
        columns[name]
        for name in ("site", "owner")
        if isinstance(columns.get(name), str) and columns[name].strip()
    )
    if len(set(mapped_headers)) != len(mapped_headers):
        raise ReportError("Chaque champ configuré doit correspondre à un en-tête CSV distinct.")

    severity_values = config.get("severity_include")
    if not isinstance(severity_values, list) or any(not isinstance(value, str) for value in severity_values):
        raise ReportError("« severity_include » doit être une liste de niveaux de sévérité.")
    severities = {value.strip().casefold() for value in severity_values}

    reports_config = config.get("reports")
    if not isinstance(reports_config, list) or len(reports_config) < 2:
        raise ReportError("« reports » doit contenir au moins deux rapports, dans l'ordre chronologique.")

    base_dir = config_path.resolve().parent
    reports: list[ScanReport] = []
    for index, item in enumerate(reports_config, start=1):
        if not isinstance(item, dict) or not isinstance(item.get("path"), str) or not item["path"].strip():
            raise ReportError(f"Le rapport numéro {index} doit définir un chemin « path » non vide.")
        for metadata in ("site", "owner", "label", "sheet"):
            if metadata in item and item[metadata] is not None and not isinstance(item[metadata], str):
                raise ReportError(f"« {metadata} » du rapport numéro {index} doit être une chaîne ou null.")
        if item.get("sheet") is not None and not item["sheet"].strip():
            raise ReportError(f"« sheet » du rapport numéro {index} doit être un nom non vide ou null.")

        source = Path(item["path"])
        if not source.is_absolute():
            source = base_dir / source
        label = (item.get("label") or source.stem).strip()
        if not label:
            raise ReportError(f"Le rapport numéro {index} doit avoir un libellé non vide.")
        reports.append(
            parse_scan_report(
                source=source,
                label=label,
                columns=columns,
                severities=severities,
                default_site=item.get("site"),
                default_owner=item.get("owner"),
                sheet_name=item.get("sheet"),
            )
        )

    output = config.get("output_html", "tenable_report.html")
    if not isinstance(output, str) or not output.strip():
        raise ReportError("« output_html » doit être un chemin non vide.")
    output_path = Path(output)
    if not output_path.is_absolute():
        output_path = base_dir / output_path
    return reports, output_path


def parse_scan_report(
    source: Path,
    label: str,
    columns: dict[str, str | None],
    severities: set[str],
    default_site: str | None = None,
    default_owner: str | None = None,
    sheet_name: str | None = None,
) -> ScanReport:
    """Lit un export CSV ou Excel, filtre les sévérités et déduplique les résultats."""
    suffix = source.suffix.casefold()
    if suffix in (".xlsx", ".xlsm"):
        return _parse_excel_report(source, label, columns, severities, default_site, default_owner, sheet_name)
    if suffix != ".csv":
        raise ReportError(
            f"Format de rapport non pris en charge pour « {source} ». Utilisez un fichier .csv, .xlsx ou .xlsm."
        )
    if sheet_name is not None:
        raise ReportError(f"L'option « sheet » ne s'applique qu'aux fichiers Excel : « {source} ».")

    try:
        stream = source.open("r", encoding="utf-8-sig", newline="")
    except OSError as exc:
        raise ReportError(f"Impossible d'ouvrir le rapport « {source} » : {exc}") from exc

    try:
        with stream:
            reader = csv.DictReader(stream)
            rows = (
                ({key: value for key, value in row.items()})
                for row in reader
            )
            findings = _parse_rows(
                source, reader.fieldnames, rows, columns, severities, default_site, default_owner, strict_rows=True
            )
    except csv.Error as exc:
        raise ReportError(f"CSV invalide dans « {source} » : {exc}") from exc
    except UnicodeError as exc:
        raise ReportError(f"Le rapport « {source} » n'est pas encodé en UTF-8.") from exc
    return ScanReport(label, source, findings)


def parse_csv_report(
    source: Path,
    label: str,
    columns: dict[str, str | None],
    severities: set[str],
    default_site: str | None = None,
    default_owner: str | None = None,
) -> ScanReport:
    """Compatibilité : lit un rapport CSV."""
    return parse_scan_report(source, label, columns, severities, default_site, default_owner)


def _parse_excel_report(
    source: Path,
    label: str,
    columns: dict[str, str | None],
    severities: set[str],
    default_site: str | None,
    default_owner: str | None,
    sheet_name: str | None,
) -> ScanReport:
    try:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise ReportError(
            "La lecture Excel nécessite openpyxl. Installez les dépendances avec "
            "« python -m pip install -r requirements-report.txt »."
        ) from exc

    try:
        workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    except (OSError, zipfile.BadZipFile, InvalidFileException) as exc:
        raise ReportError(f"Impossible de lire le fichier Excel « {source} » : {exc}") from exc
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise ReportError(
                    f"Feuille « {sheet_name} » introuvable dans « {source} ». "
                    f"Feuilles disponibles : {', '.join(workbook.sheetnames)}."
                )
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            header_cells = next(rows)
        except StopIteration as exc:
            raise ReportError(f"La feuille Excel « {worksheet.title} » de « {source} » est vide.") from exc

        headers = [_cell_text(value) for value in header_cells]
        while headers and not headers[-1]:
            headers.pop()
        def excel_rows() -> Iterator[dict[str | None, str | None]]:
            for row in rows:
                values = [_cell_text(value) for value in row]
                while len(values) > len(headers) and not values[-1]:
                    values.pop()
                yield {
                    (headers[index] if index < len(headers) else None): value
                    for index, value in enumerate(values)
                }

        data_rows = excel_rows()
        findings = _parse_rows(
            source, headers, data_rows, columns, severities, default_site, default_owner, strict_rows=False
        )
        return ScanReport(label, source, findings)
    finally:
        workbook.close()


def _parse_rows(
    source: Path,
    headers: list[str] | None,
    rows: Iterable[dict[str | None, str | None]],
    columns: dict[str, str | None],
    severities: set[str],
    default_site: str | None,
    default_owner: str | None,
    strict_rows: bool,
) -> dict[FindingKey, Finding]:
    if not headers:
        raise ReportError(f"Le rapport « {source} » est vide ou ne contient pas d'en-tête.")
    if len(set(headers)) != len(headers):
        raise ReportError(f"Le rapport « {source} » contient des en-têtes en double.")
    required_headers = [columns[name] for name in REQUIRED_COLUMNS]
    missing = [header for header in required_headers if header not in headers]
    missing.extend(
        header
        for name in ("site", "owner")
        if (header := columns.get(name)) and header not in headers
    )
    if missing:
        raise ReportError(f"Colonnes absentes du rapport « {source} » : {', '.join(dict.fromkeys(missing))}.")

    findings: dict[FindingKey, Finding] = {}
    for line_number, row in enumerate(rows, start=2):
        if not row or not any(value for value in row.values() if value is not None):
            continue
        if None in row or (strict_rows and any(value is None for value in row.values())):
            raise ReportError(f"Ligne mal formée dans « {source} », ligne {line_number}.")
        severity = _value(row, columns["severity"])
        if severity.casefold() not in severities:
            continue

        raw_ip = _value(row, columns["ip"])
        try:
            ip = ipaddress.ip_address(raw_ip).compressed
        except ValueError as exc:
            raise ReportError(
                f"Adresse IP invalide dans « {source} », ligne {line_number} : {raw_ip}"
            ) from exc

        port = _value(row, columns["port"])
        plugin_id = _value(row, columns["plugin_id"])
        cve = _value(row, columns["cve"]).upper()
        if not port or not plugin_id:
            raise ReportError(f"Port ou Plugin ID vide dans « {source} », ligne {line_number}.")
        if port.isdecimal():
            port = str(int(port))
        if plugin_id.isdecimal():
            plugin_id = str(int(plugin_id))

        site_column = columns.get("site")
        owner_column = columns.get("owner")
        site = (_value(row, site_column) if site_column else "") or default_site or UNASSIGNED
        owner = (_value(row, owner_column) if owner_column else "") or default_owner or UNASSIGNED
        finding = Finding(ip, port, plugin_id, cve, site.strip(), owner.strip(), severity)
        findings.setdefault(finding.key, finding)
    return findings


def _cell_text(value: object) -> str:
    """Convertit les valeurs Excel en texte, en évitant les suffixes numériques .0."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _value(row: dict[str | None, str | None], header: str | None) -> str:
    if header is None:
        return ""
    return (row.get(header) or "").strip()


def compare_reports(reports: Iterable[ScanReport]) -> list[Comparison]:
    """Compare chaque version à la précédente selon le quadruplet demandé."""
    ordered = list(reports)
    if len(ordered) < 2:
        raise ReportError("La comparaison nécessite au moins deux rapports.")

    comparisons: list[Comparison] = []
    previous: dict[FindingKey, Finding] = {}
    for index, report in enumerate(ordered):
        current = report.findings
        if index == 0:
            new = current
            persistent: dict[FindingKey, Finding] = {}
            fixed: dict[FindingKey, Finding] = {}
        else:
            current_keys = current.keys()
            previous_keys = previous.keys()
            new = {key: current[key] for key in current_keys - previous_keys}
            persistent = {key: current[key] for key in current_keys & previous_keys}
            fixed = {key: previous[key] for key in previous_keys - current_keys}
        comparisons.append(
            Comparison(
                report.label,
                tuple(new.values()),
                tuple(persistent.values()),
                tuple(fixed.values()),
            )
        )
        previous = current
    return comparisons


def _metric(findings: Iterable[Finding]) -> tuple[int, int]:
    values = tuple(findings)
    return len(values), len({finding.ip for finding in values})


def _risk_summary(report: ScanReport, comparison: Comparison) -> str:
    """Résume les vulnérabilités actives et nouvelles par niveau de sévérité."""
    active_by_risk: dict[str, list[Finding]] = defaultdict(list)
    new_by_risk: dict[str, list[Finding]] = defaultdict(list)
    for finding in report.findings.values():
        active_by_risk[finding.severity].append(finding)
    for finding in comparison.new:
        new_by_risk[finding.severity].append(finding)

    risks = set(active_by_risk) | set(new_by_risk)
    priority = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}
    ordered_risks = sorted(risks, key=lambda risk: (priority.get(risk.casefold(), 5), risk.casefold()))
    rows = []
    for risk in ordered_risks:
        active_count, affected_hosts = _metric(active_by_risk[risk])
        new_count, _ = _metric(new_by_risk[risk])
        rows.append(
            f"<tr><th>{_escape(risk)}</th><td>{active_count}</td>"
            f"<td>{affected_hosts}</td><td>{new_count}</td></tr>"
        )
    if not rows:
        rows.append('<tr><td colspan="4">Aucune vulnérabilité dans les sévérités incluses.</td></tr>')
    return (
        "<h4>Vulnérabilités par niveau de risque</h4><div class=\"table-wrap\"><table><thead><tr>"
        "<th>Risque</th><th>Total actif</th><th>Hôtes touchés</th><th>Nouvelles</th>"
        "</tr></thead><tbody>" + "".join(rows) + "</tbody></table></div>"
    )


def _stacked_chart_data(
    comparisons: list[Comparison],
) -> tuple[
    list[str],
    list[str],
    dict[str, list[int]],
    list[int],
    list[int],
    tuple[list[int], list[int], list[int]],
]:
    """Prépare les séries des graphiques en comptant chaque vulnérabilité ou hôte correctement."""
    labels = [comparison.label for comparison in comparisons]
    risk_order = {"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}
    risks = sorted(
        {
            finding.severity
            for comparison in comparisons
            for finding in (*comparison.new, *comparison.persistent)
        },
        key=lambda risk: (risk_order.get(risk.casefold(), 5), risk.casefold()),
    )
    risk_counts = {
        risk: [
            sum(
                finding.severity == risk
                for finding in (*comparison.new, *comparison.persistent)
            )
            for comparison in comparisons
        ]
        for risk in risks
    }
    new_hosts: list[int] = []
    persistent_hosts: list[int] = []
    shared_hosts: list[int] = []
    for comparison in comparisons:
        new_ips = {finding.ip for finding in comparison.new}
        persistent_ips = {finding.ip for finding in comparison.persistent}
        new_hosts.append(len(new_ips - persistent_ips))
        persistent_hosts.append(len(persistent_ips - new_ips))
        shared_hosts.append(len(new_ips & persistent_ips))
    return (
        labels,
        [str(risk) for risk in risks],
        risk_counts,
        [_metric(comparison.new)[0] for comparison in comparisons],
        [_metric(comparison.persistent)[0] for comparison in comparisons],
        (new_hosts, persistent_hosts, shared_hosts),
    )


def _chart_data_uri(comparisons: list[Comparison]) -> str:
    """Crée trois diagrammes empilés PNG en mémoire, sans fichier externe au HTML."""
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as pyplot
    except ImportError as exc:
        raise ReportError(
            "Matplotlib est requis pour créer le graphique. Installez les dépendances avec "
            "« python -m pip install -r requirements-report.txt »."
        ) from exc

    labels, risks, risk_counts, new_counts, persistent_counts, host_counts = _stacked_chart_data(comparisons)
    figure, axes = pyplot.subplots(3, 1, figsize=(12, 13))

    for axis, title, ylabel, series in (
        (
            axes[0],
            "Vulnérabilités actives : nouvelles + persistantes",
            "Nombre de vulnérabilités",
            (("Nouvelles", new_counts, "#e67e22"), ("Persistantes", persistent_counts, "#2980b9")),
        ),
        (
            axes[1],
            "Vulnérabilités actives par niveau de risque",
            "Nombre de vulnérabilités",
            tuple(
                (risk, risk_counts[risk], color)
                for risk, color in zip(risks, ("#8e44ad", "#c0392b", "#e67e22", "#f1c40f", "#7f8c8d"))
            ),
        ),
        (
            axes[2],
            "Hôtes touchés : nouvelles + persistantes",
            "Nombre d'hôtes uniques",
            (
                ("Nouvelles uniquement", host_counts[0], "#e67e22"),
                ("Persistantes uniquement", host_counts[1], "#2980b9"),
                ("Nouvelles et persistantes", host_counts[2], "#27ae60"),
            ),
        ),
    ):
        bottoms = [0] * len(labels)
        for display_name, counts, color in series:
            axis.bar(labels, counts, bottom=bottoms, label=display_name, color=color)
            bottoms = [base + count for base, count in zip(bottoms, counts)]
        axis.set_title(title)
        axis.set_ylabel(ylabel)
        axis.set_ylim(bottom=0)
        axis.grid(True, axis="y", linestyle=":", alpha=0.6)
        axis.tick_params(axis="x", rotation=25)
        if series:
            axis.legend()
    figure.tight_layout()
    buffer = io.BytesIO()
    figure.savefig(buffer, format="png", dpi=140, bbox_inches="tight")
    pyplot.close(figure)
    return "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode("ascii")


def render_html(reports: list[ScanReport], comparisons: list[Comparison]) -> str:
    """Construit une page autonome, avec texte échappé et image intégrée."""
    image_uri = _chart_data_uri(comparisons)
    sections: list[str] = []
    for report, comparison in zip(reports, comparisons):
        total_findings, total_hosts = _metric(report.findings.values())
        sections.append(
            f"<section><h3>{_escape(report.label)}</h3>"
            f"<p>Rapport : <code>{_escape(str(report.source))}</code> — "
            f"{total_findings} vulnérabilités actives sur {total_hosts} hôtes.</p>"
            f"{_risk_summary(report, comparison)}"
            "<div class=\"table-wrap\"><table><thead><tr><th>Catégorie</th>"
            "<th>Vulnérabilités</th><th>Hôtes</th></tr></thead><tbody>"
            f"{_summary_rows(comparison)}</tbody></table></div>"
            f"{_group_breakdown(comparison)}</section>"
        )

    return (
        "<!doctype html><html lang=\"fr\"><head><meta charset=\"utf-8\">"
        "<meta name=\"viewport\" content=\"width=device-width, initial-scale=1\">"
        "<title>Évolution des vulnérabilités Tenable</title><style>"
        "body{font:16px/1.5 system-ui,sans-serif;color:#202938;background:#f4f6f8;"
        "margin:0;padding:2rem}main{max-width:1100px;margin:auto;background:white;"
        "padding:2rem;border-radius:12px;box-shadow:0 2px 12px #17202a18}"
        "h1,h2,h3{color:#17324d}section{border-top:1px solid #dce3ea;padding:1rem 0}"
        "img{display:block;width:100%;max-width:1000px;margin:1rem auto}"
        ".table-wrap{overflow-x:auto}table{border-collapse:collapse;width:100%;margin:1rem 0}"
        "th,td{border:1px solid #dce3ea;padding:.55rem;text-align:left}"
        "thead{background:#eaf0f5}code{overflow-wrap:anywhere}.note{color:#526273}"
        "</style></head><body><main><h1>Évolution des vulnérabilités Tenable</h1>"
        "<p class=\"note\">Les rapports sont comparés de proche en proche. "
        "L'identité d'une vulnérabilité est le quadruplet IP, port, Plugin ID et CVE. "
        "Le premier rapport sert de référence : toutes ses vulnérabilités sont classées comme nouvelles. "
        "Les éléments corrigés sont attribués au site et au propriétaire du rapport précédent ; "
        "les autres catégories utilisent les métadonnées du rapport courant.</p>"
        "<h2>Tendances</h2>"
        f"<img alt=\"Trois diagrammes empilés : vulnérabilités nouvelles et persistantes, "
        f"vulnérabilités par risque et hôtes touchés\" src=\"{image_uri}\">"
        "<h2>Synthèse par version</h2>"
        "<div class=\"table-wrap\"><table><thead><tr><th>Version</th><th>Vulnérabilités actives</th>"
        "<th>Hôtes actifs</th><th>Nouvelles vulnérabilités</th><th>Hôtes nouveaux</th>"
        "<th>Persistantes</th><th>Hôtes persistants</th><th>Corrigées</th><th>Hôtes corrigés</th>"
        "</tr></thead><tbody>"
        + "".join(
            f"<tr><th>{_escape(report.label)}</th>"
            f"<td>{_metric(report.findings.values())[0]}</td><td>{_metric(report.findings.values())[1]}</td>"
            + "".join(
                f"<td>{_metric(getattr(comparison, name))[metric]}</td>"
                for name in ("new", "persistent", "fixed")
                for metric in (0, 1)
            )
            + "</tr>"
            for report, comparison in zip(reports, comparisons)
        )
        + "</tbody></table></div><h2>Détails par site et propriétaire</h2>"
        + "".join(sections)
        + "</main></body></html>"
    )


def export_excel(reports: list[ScanReport], comparisons: list[Comparison], output: Path) -> None:
    """Exporte les synthèses et les vulnérabilités classées dans un classeur Excel."""
    if output.suffix.casefold() != ".xlsx":
        raise ReportError(f"Le fichier de sortie Excel doit avoir l'extension .xlsx : « {output} ».")
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ReportError(
            "La génération Excel nécessite openpyxl. Installez les dépendances avec "
            "« python -m pip install -r requirements-report.txt »."
        ) from exc

    workbook = openpyxl.Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Synthèse"
    summary_sheet.append([
        "Scan",
        "Fichier source",
        "Vulnérabilités actives",
        "Hôtes actifs",
        "Nouvelles",
        "Hôtes nouveaux",
        "Persistantes",
        "Hôtes persistants",
        "Corrigées",
        "Hôtes corrigés",
    ])

    risk_sheet = workbook.create_sheet("Synthèse par risque")
    risk_sheet.append([
        "Scan", "Risque", "Vulnérabilités actives", "Hôtes touchés", "Nouvelles vulnérabilités",
    ])

    details_sheet = workbook.create_sheet("Détails")
    details_sheet.append([
        "Scan", "Statut", "Risque", "IP", "Port", "Plugin ID", "CVE", "Site", "Propriétaire",
    ])

    for report, comparison in zip(reports, comparisons):
        active_count, active_hosts = _metric(report.findings.values())
        summary_sheet.append([
            report.label,
            str(report.source),
            active_count,
            active_hosts,
            *_metric(comparison.new),
            *_metric(comparison.persistent),
            *_metric(comparison.fixed),
        ])

        active_by_risk: dict[str, list[Finding]] = defaultdict(list)
        new_by_risk: dict[str, list[Finding]] = defaultdict(list)
        for finding in report.findings.values():
            active_by_risk[finding.severity].append(finding)
        for finding in comparison.new:
            new_by_risk[finding.severity].append(finding)
        for risk in sorted(
            set(active_by_risk) | set(new_by_risk),
            key=lambda value: ({"critical": 0, "high": 1, "medium": 2, "low": 3, "info": 4}.get(value.casefold(), 5),
                               value.casefold()),
        ):
            active_risk_count, affected_hosts = _metric(active_by_risk[risk])
            new_risk_count, _ = _metric(new_by_risk[risk])
            risk_sheet.append([report.label, risk, active_risk_count, affected_hosts, new_risk_count])

        for status, findings in (
            ("Nouvelle", comparison.new),
            ("Persistante", comparison.persistent),
            ("Corrigée", comparison.fixed),
        ):
            for finding in findings:
                details_sheet.append([
                    report.label,
                    status,
                    finding.severity,
                    finding.ip,
                    finding.port,
                    finding.plugin_id,
                    finding.cve,
                    finding.site,
                    finding.owner,
                ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="17324D")
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = width

    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    except OSError as exc:
        raise ReportError(f"Impossible d'écrire le fichier Excel « {output} » : {exc}") from exc
    finally:
        workbook.close()


def _summary_rows(comparison: Comparison) -> str:
    rows = []
    for title, name in (("Nouvelles", "new"), ("Persistantes", "persistent"), ("Corrigées", "fixed")):
        count, hosts = _metric(getattr(comparison, name))
        rows.append(f"<tr><th>{title}</th><td>{count}</td><td>{hosts}</td></tr>")
    return "".join(rows)


def _group_breakdown(comparison: Comparison) -> str:
    grouped: dict[tuple[str, str], dict[str, list[Finding]]] = defaultdict(
        lambda: {"new": [], "persistent": [], "fixed": []}
    )
    for category in ("new", "persistent", "fixed"):
        for finding in getattr(comparison, category):
            grouped[(finding.site, finding.owner)][category].append(finding)
    if not grouped:
        return "<p>Aucune vulnérabilité classée pour cette version.</p>"
    rows = []
    for (site, owner), values in sorted(grouped.items()):
        metrics = [_metric(values[name]) for name in ("new", "persistent", "fixed")]
        rows.append(
            "<tr>"
            f"<td>{_escape(site)}</td><td>{_escape(owner)}</td>"
            + "".join(f"<td>{count}</td><td>{hosts}</td>" for count, hosts in metrics)
            + "</tr>"
        )
    return (
        "<h4>Détail par site / propriétaire</h4><div class=\"table-wrap\"><table><thead><tr>"
        "<th>Site</th><th>Propriétaire</th><th>Nouvelles</th><th>Hôtes nouveaux</th>"
        "<th>Persistantes</th><th>Hôtes persistants</th><th>Corrigées</th><th>Hôtes corrigés</th>"
        "</tr></thead><tbody>" + "".join(rows) + "</tbody></table></div>"
    )


def _escape(value: str) -> str:
    return html.escape(value, quote=True)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Compare des exports CSV Tenable et génère un rapport HTML autonome.")
    parser.add_argument("--config", type=Path, default=Path("config.json"), help="Fichier de configuration JSON.")
    parser.add_argument("--output", type=Path, help="Remplace le chemin output_html configuré.")
    parser.add_argument("--excel-output", type=Path, help="Chemin du classeur Excel de résultats.")
    args = parser.parse_args(argv)
    try:
        reports, output = load_configuration(args.config)
        print("Fichiers de scan utilisés :")
        for index, report in enumerate(reports, start=1):
            print(f"  {index}. {report.label} : {report.source}")
        if args.output:
            output = args.output if args.output.is_absolute() else Path.cwd() / args.output
        excel_output = args.excel_output
        if excel_output is None:
            excel_output = output.with_suffix(".xlsx")
        elif not excel_output.is_absolute():
            excel_output = Path.cwd() / excel_output
        if excel_output.resolve() in {report.source.resolve() for report in reports}:
            raise ReportError("Le fichier Excel de sortie ne peut pas remplacer un fichier de scan source.")
        comparisons = compare_reports(reports)
        document = render_html(reports, comparisons)
        export_excel(reports, comparisons, excel_output)
        output.parent.mkdir(parents=True, exist_ok=True)
        output.write_text(document, encoding="utf-8")
    except (ReportError, OSError) as exc:
        print(f"Erreur : {exc}", file=sys.stderr)
        return 1
    print(f"Rapport généré : {output}")
    print(f"Classeur Excel généré : {excel_output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

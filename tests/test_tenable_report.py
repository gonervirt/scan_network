import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

import tenable_report


DATA_DIR = Path(__file__).parent / "data"


class CsvParsingTests(unittest.TestCase):
    def test_reads_configurable_headers_filters_severity_and_deduplicates(self):
        reports, _ = tenable_report.load_configuration(DATA_DIR / "config.json")
        first = reports[0]
        self.assertEqual(len(first.findings), 2)
        finding = next(iter(first.findings.values()))
        self.assertEqual(finding.key, ("10.0.0.1", "443", "1001", "CVE-2024-0001"))
        self.assertEqual((finding.site, finding.owner), ("Paris", "Ops"))
        self.assertEqual(finding.severity, "High")

    def test_report_metadata_is_used_when_optional_columns_are_absent(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "scan.csv"
            source.write_text(
                "Address,Port,Plugin,CVE,Severity\n"
                "10.1.2.3,443,0007,cve-2024-0100,High\n",
                encoding="utf-8",
            )
            report = tenable_report.parse_csv_report(
                source,
                "Scan",
                {"ip": "Address", "port": "Port", "plugin_id": "Plugin", "cve": "CVE", "severity": "Severity"},
                {"high"},
                default_site="Site A",
                default_owner="Owner B",
            )
        finding = next(iter(report.findings.values()))
        self.assertEqual(finding.key, ("10.1.2.3", "443", "7", "CVE-2024-0100"))
        self.assertEqual((finding.site, finding.owner), ("Site A", "Owner B"))

    def test_missing_required_header_is_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "scan.csv"
            source.write_text("IP,Port\n10.1.1.1,443\n", encoding="utf-8")
            with self.assertRaisesRegex(tenable_report.ReportError, "Colonnes absentes"):
                tenable_report.parse_csv_report(
                    source,
                    "Scan",
                    {"ip": "IP", "port": "Port", "plugin_id": "Plugin", "cve": "CVE", "severity": "Severity"},
                    {"high"},
                )


class ExcelParsingTests(unittest.TestCase):
    def test_reads_xlsx_with_configurable_headers_and_selected_sheet(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "scan.xlsx"
            workbook = openpyxl.Workbook()
            default_sheet = workbook.active
            default_sheet.title = "Summary"
            default_sheet.append(["This is not the scan data"])
            findings_sheet = workbook.create_sheet("Findings")
            findings_sheet.append(["Address", "Port", "Plugin", "CVE", "Severity"])
            findings_sheet.append(["10.1.2.3", 443, 7, "cve-2024-0100", "High"])
            findings_sheet.append(["10.1.2.4", 80, 8, "CVE-2024-0101", "Low"])
            workbook.save(source)
            workbook.close()

            report = tenable_report.parse_scan_report(
                source,
                "Excel scan",
                {"ip": "Address", "port": "Port", "plugin_id": "Plugin", "cve": "CVE", "severity": "Severity"},
                {"high"},
                sheet_name="Findings",
            )

        self.assertEqual(len(report.findings), 1)
        finding = next(iter(report.findings.values()))
        self.assertEqual(finding.key, ("10.1.2.3", "443", "7", "CVE-2024-0100"))

    def test_configuration_can_compare_excel_and_csv_reports(self):
        columns = {"ip": "IP", "port": "Port", "plugin_id": "Plugin", "cve": "CVE", "severity": "Severity"}
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(list(columns.values()))
            sheet.append(["10.1.2.3", 443, 7, "CVE-2024-0100", "High"])
            workbook.save(root / "first.xlsx")
            workbook.close()
            (root / "second.csv").write_text(
                "IP,Port,Plugin,CVE,Severity\n10.1.2.3,443,7,CVE-2024-0100,High\n",
                encoding="utf-8",
            )
            config_path = root / "config.json"
            config_path.write_text(
                json.dumps(
                    {
                        "reports": [{"path": "first.xlsx"}, {"path": "second.csv"}],
                        "columns": columns,
                        "severity_include": ["High"],
                    }
                ),
                encoding="utf-8",
            )

            reports, _ = tenable_report.load_configuration(config_path)

        self.assertEqual([len(report.findings) for report in reports], [1, 1])

    def test_unsupported_excel_extension_has_clear_error(self):
        with self.assertRaisesRegex(tenable_report.ReportError, "xlsx ou .xlsm"):
            tenable_report.parse_scan_report(
                Path("scan.xls"),
                "Excel scan",
                {"ip": "IP", "port": "Port", "plugin_id": "Plugin", "cve": "CVE", "severity": "Severity"},
                {"high"},
            )


class ComparisonTests(unittest.TestCase):
    def test_stacked_chart_data_counts_active_risks_and_unique_host_groups(self):
        persistent = tenable_report.Finding(
            "10.0.0.1", "443", "1001", "CVE-2024-0001", "Site", "Owner", "High"
        )
        new_shared_host = tenable_report.Finding(
            "10.0.0.1", "80", "1002", "CVE-2024-0002", "Site", "Owner", "Critical"
        )
        new_host = tenable_report.Finding(
            "10.0.0.2", "22", "1003", "CVE-2024-0003", "Site", "Owner", "High"
        )
        comparison = tenable_report.Comparison(
            "Scan 2",
            (new_shared_host, new_host),
            (persistent,),
            (),
        )

        labels, risks, risk_counts, new_counts, persistent_counts, host_counts = (
            tenable_report._stacked_chart_data([comparison])
        )

        self.assertEqual(labels, ["Scan 2"])
        self.assertEqual(risks, ["Critical", "High"])
        self.assertEqual(risk_counts, {"Critical": [1], "High": [2]})
        self.assertEqual(new_counts, [2])
        self.assertEqual(persistent_counts, [1])
        self.assertEqual(host_counts, ([1], [0], [1]))
        self.assertEqual(sum(series[0] for series in host_counts), 2)

    def test_excel_export_contains_summaries_and_classified_findings(self):
        reports, _ = tenable_report.load_configuration(DATA_DIR / "config.json")
        comparisons = tenable_report.compare_reports(reports)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "results.xlsx"

            tenable_report.export_excel(reports, comparisons, output)

            workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["Synthèse", "Synthèse par risque", "Détails"])
            summary = workbook["Synthèse"]
            self.assertEqual(summary["A2"].value, "January 2024")
            self.assertEqual(summary["C2"].value, 2)
            risk_rows = list(workbook["Synthèse par risque"].iter_rows(values_only=True))
            self.assertIn(("January 2024", "High", 1, 1, 1), risk_rows)
            detail_rows = list(workbook["Détails"].iter_rows(values_only=True))
            self.assertIn(
                ("February 2024", "Nouvelle", "Medium", "10.0.0.1", "80", "1004",
                 "CVE-2024-0004", "Lyon", "Sec"),
                detail_rows,
            )
            self.assertIn(
                ("February 2024", "Corrigée", "Critical", "10.0.0.2", "22", "1002",
                 "CVE-2024-0002", "Paris", "Ops"),
                detail_rows,
            )
            workbook.close()

    def test_excel_export_rejects_non_xlsx_extension(self):
        with self.assertRaisesRegex(tenable_report.ReportError, "extension .xlsx"):
            tenable_report.export_excel([], [], Path("results.csv"))

    def test_risk_summary_reports_active_hosts_and_new_vulnerabilities(self):
        findings = [
            tenable_report.Finding("10.0.0.1", "443", "1001", "CVE-2024-0001", "Site", "Owner", "High"),
            tenable_report.Finding("10.0.0.1", "80", "1002", "CVE-2024-0002", "Site", "Owner", "High"),
            tenable_report.Finding("10.0.0.2", "22", "1003", "CVE-2024-0003", "Site", "Owner", "Medium"),
        ]
        current = tenable_report.ScanReport(
            "Scan courant",
            Path("scan.csv"),
            {finding.key: finding for finding in findings},
        )
        previous = tenable_report.ScanReport("Scan précédent", Path("previous.csv"), {})
        comparison = tenable_report.compare_reports([previous, current])[1]

        summary = tenable_report._risk_summary(current, comparison)

        self.assertIn("<tr><th>High</th><td>2</td><td>1</td><td>2</td></tr>", summary)
        self.assertIn("<tr><th>Medium</th><td>1</td><td>1</td><td>1</td></tr>", summary)

    def test_sequential_deltas_have_expected_counts_and_unique_hosts(self):
        def scan(label, entries):
            findings = {
                finding.key: finding
                for finding in (
                    tenable_report.Finding(ip, port, plugin, cve, "Site", "Owner")
                    for ip, port, plugin, cve in entries
                )
            }
            return tenable_report.ScanReport(label, Path(label), findings)

        reports = [
            scan("janvier", [
                ("10.0.0.1", "443", "1001", "CVE-2024-0001"),
                ("10.0.0.2", "22", "1002", "CVE-2024-0002"),
            ]),
            scan("février", [
                ("10.0.0.1", "443", "1001", "CVE-2024-0001"),
                ("10.0.0.1", "80", "1004", "CVE-2024-0004"),
            ]),
            scan("mars", [
                ("10.0.0.1", "80", "1004", "CVE-2024-0004"),
                ("10.0.0.5", "443", "1006", "CVE-2024-0006"),
            ]),
        ]
        jan, feb, mar = tenable_report.compare_reports(reports)

        self.assertEqual(tenable_report._metric(jan.new), (2, 2))
        self.assertEqual(tenable_report._metric(jan.persistent), (0, 0))
        self.assertEqual(tenable_report._metric(jan.fixed), (0, 0))
        self.assertEqual(tenable_report._metric(feb.new), (1, 1))
        self.assertEqual(tenable_report._metric(feb.persistent), (1, 1))
        self.assertEqual(tenable_report._metric(feb.fixed), (1, 1))
        self.assertEqual(tenable_report._metric(mar.new), (1, 1))
        self.assertEqual(tenable_report._metric(mar.persistent), (1, 1))
        self.assertEqual(tenable_report._metric(mar.fixed), (1, 1))

    def test_html_is_self_contained_and_escapes_metadata(self):
        reports, _ = tenable_report.load_configuration(DATA_DIR / "config.json")
        first = reports[0]
        altered = tenable_report.ScanReport("<script>alert(1)</script>", first.source, first.findings)
        document = tenable_report.render_html(
            [altered, reports[1], reports[2]],
            tenable_report.compare_reports([altered, reports[1], reports[2]]),
        )
        self.assertIn("data:image/png;base64,", document)
        self.assertIn('<html lang="en">', document)
        self.assertIn("<h2>Trends</h2>", document)
        self.assertIn("<th>Active vulnerabilities</th>", document)
        self.assertIn("Breakdown by site / owner", document)
        self.assertNotIn("Vulnérabilités actives", document)
        self.assertIn("&lt;script&gt;", document)
        self.assertNotIn("<script>alert(1)</script>", document)
        self.assertNotIn('src="tenable_report', document)


if __name__ == "__main__":
    unittest.main()
